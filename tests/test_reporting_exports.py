"""Report assembly, Markdown rendering, and export tests."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from salespulse.analytics import AnalyticsService
from salespulse.exports import (
    export_bundle_to_csv,
    export_bundle_to_excel,
    export_quality_csv,
    sanitize_frame,
    sanitize_spreadsheet_value,
)
from salespulse.reporting import (
    format_report_text,
    generate_report,
    write_markdown_report,
)

EXPECTED_SECTIONS = {
    "Overview",
    "Order Status",
    "Daily Revenue",
    "Monthly Revenue",
    "Monthly Revenue MoM",
    "Product Performance",
    "Category Performance",
    "Top Products by Revenue",
    "Customer New vs Returning",
    "Customer Summary",
    "Customer Frequency",
    "Top Customers",
    "Regional Performance",
    "Cancellation & Return Rates",
    "Monthly Rankings",
    "Data Quality",
    "Import History",
    "Rejected Records",
}


@pytest.fixture(scope="module")
def bundle(clean_engine):
    service = AnalyticsService(clean_engine)
    start, end = service.date_bounds()
    return generate_report(clean_engine, start, end)


def test_bundle_contains_all_sections_and_metadata(bundle) -> None:
    assert set(bundle.sections) == EXPECTED_SECTIONS
    assert bundle.metadata["title"] == "SalesPulse Sales Performance Report"
    assert bundle.metadata["synthetic_notice"]
    assert "approved refunds" in bundle.metadata["revenue_definition"]
    assert bundle.generated_at.endswith("Z")
    assert bundle.metadata["quality_summary"]["critical_failures"] == 0

    overview = bundle.section("Overview").iloc[0]
    assert int(overview["completed_orders"]) > 0
    assert float(overview["net_revenue"]) > 0
    assert bundle.section("Data Quality")["status"].eq("pass").all()


def test_report_text_is_human_readable(bundle) -> None:
    text = format_report_text(bundle, max_rows=3)
    assert "SalesPulse Sales Performance Report" in text
    assert "Synthetic demonstration data" in text
    assert "Net revenue" in text
    assert "KPI summary" in text
    assert "more rows available in exports" in text


def test_markdown_report_is_written(bundle, tmp_path: Path) -> None:
    path = write_markdown_report(bundle, tmp_path / "report.md")
    text = path.read_text(encoding="utf-8")
    assert text.startswith("# SalesPulse Sales Performance Report")
    assert "Synthetic demonstration data" in text
    assert "## KPI summary" in text
    assert "## Monthly revenue" in text
    assert "## Data quality" in text
    assert "nan" not in text.lower()
    assert "None" not in text


def test_excel_export_has_readme_and_valid_sheets(bundle, tmp_path: Path) -> None:
    path = export_bundle_to_excel(bundle, tmp_path / "report.xlsx")
    assert path.exists() and path.stat().st_size > 10_000
    workbook = load_workbook(path, read_only=True)
    names = workbook.sheetnames
    assert names[0] == "Read Me"
    assert "Overview" in names
    assert "Monthly Revenue" in names
    assert len(names) == len(set(names))
    assert all(len(name) <= 31 for name in names)
    readme = workbook["Read Me"]
    values = [cell.value for row in readme.iter_rows() for cell in row]
    assert any(value == "SalesPulse — Sales Performance Report" for value in values)
    assert any("Synthetic" in str(value) for value in values if value)
    workbook.close()


def test_csv_export_has_manifest_and_all_sections(bundle, tmp_path: Path) -> None:
    directory = export_bundle_to_csv(bundle, tmp_path / "csv")
    manifest_path = directory / "report_manifest.json"
    assert manifest_path.exists()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["title"] == bundle.title
    assert manifest["date_from"] == bundle.date_from
    assert manifest["synthetic_notice"]
    assert len(manifest["files"]) == len(bundle.sections)
    for record in manifest["files"]:
        path = directory / record["filename"]
        assert path.exists(), record["section"]
        assert record["rows"] == len(bundle.section(record["section"]))
        assert len(record["sha256"]) == 64
        assert record["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert not any(" " in record["filename"] for record in manifest["files"])


def test_quality_csv_export(bundle, tmp_path: Path) -> None:
    path = export_quality_csv(bundle, tmp_path / "quality.csv")
    frame = pd.read_csv(path)
    assert len(frame) == 25
    assert set(frame["status"]) == {"pass"}


def test_empty_range_exports_do_not_crash(clean_engine, tmp_path: Path) -> None:
    bundle = generate_report(clean_engine, "2035-01-01", "2035-01-31")
    xlsx = export_bundle_to_excel(bundle, tmp_path / "empty.xlsx")
    assert xlsx.exists()
    csv_dir = export_bundle_to_csv(bundle, tmp_path / "empty_csv")
    assert (csv_dir / "report_manifest.json").exists()
    markdown = write_markdown_report(bundle, tmp_path / "empty.md")
    assert markdown.exists()
    assert bundle.section("Overview").iloc[0]["completed_orders"] == 0


def test_spreadsheet_formula_injection_is_neutralized() -> None:
    assert sanitize_spreadsheet_value("=1+1") == "'=1+1"
    assert sanitize_spreadsheet_value("+SUM(A1)") == "'+SUM(A1)"
    assert sanitize_spreadsheet_value("@cmd") == "'@cmd"
    assert sanitize_spreadsheet_value(12.5) == 12.5
    frame = sanitize_frame(pd.DataFrame({"name": ["=evil()", "safe"], "value": [1, 2]}))
    assert frame.loc[0, "name"] == "'=evil()"
    assert frame.loc[0, "value"] == 1


def test_category_filter_is_applied(clean_engine, bundle) -> None:
    category = AnalyticsService(clean_engine).available_categories()[0]
    filtered = generate_report(
        clean_engine,
        bundle.date_from,
        bundle.date_to,
        category=category,
    )
    assert filtered.metadata["category_filter"] == category
    products = filtered.section("Product Performance")
    assert set(products["category"]) == {category}
    top = filtered.section("Top Products by Revenue")
    assert len(top) <= 10
